"""Builds a small .docx with known violations, used by unit and E2E tests."""

from pathlib import Path

from docx import Document
from docx.shared import Pt

INTRO = "The Falcon radar is a TOP SECRET prototype developed for coastal surveillance."
PERFORMANCE = "The maximum detection range is 50 km and the tracking accuracy is 3 m."
WRONG_FONT = "This sentence uses the wrong font."
LONG_SENTENCE = (
    "This deliberately long sentence keeps going and going with many extra words so that "
    "it clearly exceeds the configured maximum number of words that a single sentence is "
    "allowed to contain in a well written technical document."
)


def build_sample_docx(path: Path) -> Path:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    doc.add_heading("System Overview", level=0)
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph(INTRO)

    doc.add_heading("Performance", level=1)
    doc.add_paragraph(PERFORMANCE)

    p = doc.add_paragraph("Formatting check: ")
    run = p.add_run(WRONG_FONT)
    run.font.name = "Comic Sans MS"

    doc.add_paragraph(LONG_SENTENCE)

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Update rate"
    table.cell(1, 1).text = "2 seconds"

    # Merged cell: must be emitted once, at its first grid position
    merged = doc.add_table(rows=1, cols=3)
    merged.cell(0, 0).merge(merged.cell(0, 1)).text = "Merged heading"
    merged.cell(0, 2).text = "Last"

    doc.save(str(path))
    return path

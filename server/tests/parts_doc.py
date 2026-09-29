"""A .docx with every part the parser reads: headers/footers, footnotes, endnotes, text boxes,
nested tables, hyperlinks and field results. python-docx can't author most of these, so the XML is
written directly."""

from pathlib import Path

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = (
    f'xmlns:w="{W}" xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
)

HEADER = "Classified header: TOP SECRET"
FIRST_PAGE_HEADER = "Cover page header"
FOOTER = "Footer contact: ops@example.com"
FOOTNOTE_1 = "Measured range was 85 km in trials (TOP SECRET)."
FOOTNOTE_2 = "Second note."
ENDNOTE = "Endnote with TOP SECRET marking."
TEXT_BOX = "Text box says TOP SECRET"
NESTED_CELL = "Nested cell TOP SECRET"
LINK_TEXT = "the design portal"
CAPTION = "Figure 1: Block diagram"


def _run(text: str) -> str:
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _notes(tag: str, bodies: list[str]) -> bytes:
    separators = (
        f'<w:{tag} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{tag}>'
        f'<w:{tag} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{tag}>'
    )
    notes = "".join(
        f'<w:{tag} w:id="{i}"><w:p><w:r><w:{tag}Ref/></w:r>{_run(" " + body)}</w:p></w:{tag}>'
        for i, body in enumerate(bodies, start=1)
    )
    return (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:{tag}s {NS}>{separators}{notes}</w:{tag}s>'.encode()
    )


def build_parts_docx(path: Path) -> Path:
    doc = Document()

    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.header.paragraphs[0].text = HEADER
    section.first_page_header.paragraphs[0].text = FIRST_PAGE_HEADER
    section.footer.paragraphs[0].text = FOOTER

    doc.add_heading("Design", level=1)
    doc.add_paragraph("Body text before the box.")

    # Text box stored twice (DrawingML choice + VML fallback), as Word writes it
    box = doc.add_paragraph("Paragraph with a box.")
    box._p.append(
        parse_xml(
            f"<w:r {NS}><mc:AlternateContent>"
            f'<mc:Choice Requires="wps"><w:drawing><wps:txbx><w:txbxContent><w:p>{_run(TEXT_BOX)}</w:p>'
            "</w:txbxContent></wps:txbx></w:drawing></mc:Choice>"
            f"<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p>{_run(TEXT_BOX)}</w:p>"
            "</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>"
            "</mc:AlternateContent></w:r>"
        )
    )

    # Hyperlink and a footnote reference
    link = doc.add_paragraph("See ")
    link._p.append(parse_xml(f'<w:hyperlink {NS} w:anchor="portal">{_run(LINK_TEXT)}</w:hyperlink>'))
    link._p.append(parse_xml(f'<w:r {NS}><w:footnoteReference w:id="1"/></w:r>'))
    link._p.append(parse_xml(f"<w:r {NS}><w:t> for details.</w:t></w:r>"))

    # Caption whose number is a SEQ field result
    caption = doc.add_paragraph("Figure ", style="Caption")
    caption._p.append(parse_xml(f'<w:fldSimple {NS} w:instr=" SEQ Figure \\* ARABIC ">{_run("1")}</w:fldSimple>'))
    caption._p.append(parse_xml(f"<w:r {NS}><w:t>: Block diagram</w:t></w:r>"))
    doc.add_paragraph("As shown in Figure 1 and Figure 2.")

    # Table with a nested table
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Outer cell"
    inner = table.cell(0, 1).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = NESTED_CELL

    # Second section: its header/footer are linked to the first one, so they must not repeat
    doc.add_section()
    doc.add_paragraph("Second section body.")

    package = doc.part.package
    for name, ct, rt, xml in (
        ("footnotes", CT.WML_FOOTNOTES, RT.FOOTNOTES, _notes("footnote", [FOOTNOTE_1, FOOTNOTE_2])),
        ("endnotes", CT.WML_ENDNOTES, RT.ENDNOTES, _notes("endnote", [ENDNOTE])),
    ):
        doc.part.relate_to(Part(PackURI(f"/word/{name}.xml"), ct, xml, package), rt)

    doc.save(str(path))
    return path

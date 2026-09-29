"""Run the AI-rule evaluation.

    python -m eval                 # plan only: documents, labels and the number of API calls (free)
    python -m eval --run --yes     # call the real LLM and score the answers (spends quota)

Options: --docs name1,name2 to run a subset; --model to override LLM_MODEL; --out DIR for the JSON result.
"""

import argparse
import io
import json
import sys
import time
from collections import Counter
from dataclasses import asdict
from datetime import UTC, datetime
from pathlib import Path

from docx import Document

from app.config import get_settings
from app.docx_parser import parse_docx
from app.llm import LlmError, get_llm_provider
from app.rules.llm_check import check_llm_rules, count_llm_calls
from app.schemas import Block

from .dataset import DOCS, RULES, EvalDoc
from .scoring import DocResult, RuleScore, locate_labels, score_document, totals


def build_blocks(doc: EvalDoc) -> list[Block]:
    """Build the document as a real .docx and parse it with the production parser."""
    d = Document()
    d.add_heading(doc.title, level=0)
    for item in doc.items:
        if isinstance(item, str):
            d.add_paragraph(item)
        elif item[0] == "heading" and isinstance(item[1], str):
            d.add_heading(item[1], level=1)
        elif item[0] == "table" and isinstance(item[1], list):
            rows = item[1]
            table = d.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, text in enumerate(row):
                    table.cell(r, c).text = text
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return parse_docx(buf)


def _pct(value: float | None) -> str:
    return "  n/a" if value is None else f"{value * 100:5.1f}%"


def print_plan(docs: list[EvalDoc]) -> int:
    settings = get_settings()
    calls = 0
    print(f"Model: {settings.llm_model}   rules: {', '.join(r.name for r in RULES)}\n")
    print(f"{'document':24} {'lang':4} {'expected':>8} {'optional':>8} {'traps':>5} {'calls':>5}")
    for doc in docs:
        blocks = build_blocks(doc)
        labels = locate_labels(doc, blocks)  # also validates that every label occurs in the text
        n = count_llm_calls(blocks)
        calls += n
        count = Counter(lb.kind for lb in labels)
        print(f"{doc.name:24} {doc.language:4} {count['expected']:>8} {count['optional']:>8} {count['trap']:>5} {n:>5}")
    worst = calls * (settings.llm_max_retries + 1) * len(settings.llm_models)
    print(f"\nAPI calls: {calls} (worst case with retries/fallbacks: {worst})")
    return calls


def print_report(scores: dict[int, RuleScore], results: list[DocResult]) -> None:
    names = {r.id: r.name for r in RULES}
    print(f"\n{'rule':34} {'recall':>7} {'precision':>9} {'caught':>9} {'FP':>4}")
    for rule_id, s in sorted(scores.items()):
        row = f"{names[rule_id]:34} {_pct(s.recall):>7} {_pct(s.precision):>9}"
        print(f"{row} {s.caught:>4}/{s.expected:<4} {s.false_positives:>4}")
    t = totals(scores)
    print(
        f"{'TOTAL':34} {_pct(t.recall):>7} {_pct(t.precision):>9} {t.caught:>4}/{t.expected:<4} {t.false_positives:>4}"
    )

    for r in results:
        if r.missed or r.false_positives:
            print(f"\n{r.name}:")
            for lb in r.missed:
                print(f"  MISSED  [{names[lb.rule_id] if lb.rule_id else '-'}] {lb.quote}")
            for rule_id, text in r.false_positives:
                trap = " (trap)" if (rule_id, text) in r.trap_hits else ""
                print(f"  FALSE+  [{names[rule_id]}] {text}{trap}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m eval", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--run", action="store_true", help="call the real LLM (spends quota)")
    parser.add_argument("--yes", action="store_true", help="don't ask for confirmation before calling the LLM")
    parser.add_argument("--docs", help="comma-separated document names (default: all)")
    parser.add_argument("--model", help="override LLM_MODEL")
    parser.add_argument("--out", default="eval/results", help="directory for the JSON result")
    args = parser.parse_args(argv)

    settings = get_settings()
    if args.model:
        settings.llm_model = args.model
    docs = DOCS
    if args.docs:
        wanted = set(args.docs.split(","))
        docs = [d for d in DOCS if d.name in wanted]
        if unknown := wanted - {d.name for d in docs}:
            parser.error(f"unknown documents: {', '.join(sorted(unknown))}")

    calls = print_plan(docs)
    if not args.run:
        print("\nPlan only. Add --run to call the LLM.")
        return 0
    if not args.yes:
        if not sys.stdin.isatty():
            parser.error("--run needs --yes when not running in a terminal")
        if input(f"Make {calls} LLM calls with {settings.llm_model}? [y/N] ").strip().lower() != "y":
            print("Cancelled.")
            return 1

    provider = get_llm_provider()
    scores: dict[int, RuleScore] = {}
    results: list[DocResult] = []
    started = time.monotonic()
    for doc in docs:
        blocks = build_blocks(doc)
        try:
            findings = check_llm_rules(provider, RULES, blocks)
        except LlmError as exc:
            print(f"\n{doc.name}: LLM error, stopping: {exc}")
            return 2
        results.append(score_document(doc, blocks, findings, scores))
    elapsed = time.monotonic() - started

    print_report(scores, results)
    t = totals(scores)
    print(f"\n{len(docs)} documents, {calls} calls, {elapsed:.0f} s")

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
    path = out / f"{stamp}-{settings.llm_model}.json"
    path.write_text(
        json.dumps(
            {
                "model": settings.llm_model,
                "documents": [d.name for d in docs],
                "calls": calls,
                "seconds": round(elapsed, 1),
                "total": {**asdict(t), "recall": t.recall, "precision": t.precision},
                "rules": {
                    str(k): {**asdict(v), "recall": v.recall, "precision": v.precision} for k, v in scores.items()
                },
                "misses": [{"doc": r.name, "quote": lb.quote, "rule": lb.rule_id} for r in results for lb in r.missed],
                "false_positives": [
                    {"doc": r.name, "rule": k, "text": t} for r in results for k, t in r.false_positives
                ],
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"Saved {path}")
    return 0
